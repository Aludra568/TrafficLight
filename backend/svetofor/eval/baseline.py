"""Базовая линия: «склеили текст файлов, спросили модель» — для честного сравнения на эталоне.

Так устроено типичное решение: текст всех документов (обрезанный по контексту модели) и список
рисков — в один промпт, ответ — коды рисков с цитатами, без проверки цитат, грамматик и правил.
На том же наборе считаются полнота и точность по правилам и доля цитат, которых нет в документах
(«выдуманные цитаты»). Сравнение с нашим конвейером — по одинаковому эталону `expected.yaml`.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml
from pydantic import BaseModel, Field

from svetofor.eval.bootstrap import bootstrap_ci
from svetofor.eval.metrics import PackageResult, PairResult, precision, recall
from svetofor.eval.run import plural
from svetofor.llm.errors import LLMError
from svetofor.llm.gateway import LLMGateway
from svetofor.llm.prompting import render
from svetofor.pipeline.chunking import CHARS_PER_TOKEN
from svetofor.rules.model import Rule
from svetofor.rules.verify import QuoteVerifier, TextSource

BASELINE_TEMPLATE = "baseline.v1"
BASELINE_TASK = "baseline"


class BaselineRisk(BaseModel):
    rule_id: str
    quote: str | None = Field(None, max_length=600)


class BaselineResult(BaseModel):
    risks: list[BaselineRisk]


@dataclass(frozen=True)
class _Document:
    name: str
    text: str


def read_text(path: Path) -> str:
    """Текст файла «как у всех»: DOCX — абзацы python-docx, PDF — текстовый слой PyMuPDF."""
    if path.suffix.lower() == ".docx":
        import docx

        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if path.suffix.lower() == ".pdf":
        from svetofor.pipeline.pdf import open_pdf

        with open_pdf(path) as doc:
            return "\n".join(page.get_text("text") for page in doc)
    return ""


def _documents(folder: Path, budget_chars: int) -> list[_Document]:
    """Документы пакета, обрезанные по общему бюджету — как поступает наивное решение."""
    docs = [
        _Document(p.name, read_text(p))
        for p in sorted(folder.iterdir())
        if p.suffix.lower() in (".docx", ".pdf")
    ]
    left, result = budget_chars, []
    for doc in docs:
        text = doc.text[: max(0, left)]
        left -= len(text)
        result.append(_Document(doc.name, text))
    return result


@dataclass
class BaselinePackage:
    result: PackageResult
    quotes_total: int = 0
    quotes_found: int = 0
    error: str | None = None


def evaluate_baseline_package(
    gateway: LLMGateway, folder: Path, rules: list[Rule], context_tokens: int
) -> BaselinePackage:
    data = yaml.safe_load((folder / "expected.yaml").read_text(encoding="utf-8"))
    documents = _documents(folder, int(context_tokens * 0.6 * CHARS_PER_TOKEN))
    prompt = render(BASELINE_TEMPLATE, rules=rules, documents=documents)
    package = BaselinePackage(PackageResult(data["package"]))
    try:
        answer = gateway.complete_json(
            task=BASELINE_TASK,
            model_role="main",
            system=prompt.system,
            user=prompt.user,
            schema=BaselineResult,
            prompt_version=prompt.version,
            max_tokens=3000,
        )
        risks = answer.value.risks
    except LLMError as exc:
        package.error, package.result.failed, risks = exc.message, True, []
    predicted = {r.rule_id.strip() for r in risks}
    for item in data["expected"]:
        package.result.pairs.append(
            PairResult(item["rule"], item["present"], item["rule"] in predicted)
        )
    sources = [TextSource(d.name, d.text) for d in _documents(folder, 10**9)]
    verifier = QuoteVerifier()
    for risk in risks:
        if risk.quote:
            package.quotes_total += 1
            package.quotes_found += verifier.verify(risk.quote, sources) is not None
    return package


def _quotes_ok(packages: list[BaselinePackage]) -> float | None:
    total = sum(p.quotes_total for p in packages)
    return sum(p.quotes_found for p in packages) / total if total else None


def run_baseline(
    gateway: LLMGateway,
    set_dir: Path,
    rules: list[Rule],
    results_dir: Path,
    *,
    model: str | None,
    context_tokens: int = 32768,
    limit: int | None = None,
) -> dict[str, Any]:
    folders = sorted(p for p in set_dir.iterdir() if (p / "expected.yaml").is_file())[:limit]
    if not folders:
        raise FileNotFoundError(f"В {set_dir} нет пакетов с expected.yaml")
    packages = [evaluate_baseline_package(gateway, f, rules, context_tokens) for f in folders]
    results = [p.result for p in packages]
    report: dict[str, Any] = {
        "dataset": set_dir.name,
        "created_at": datetime.now(UTC).isoformat(timespec="seconds"),
        "model": model,
        "packages": len(packages),
        "approach": "baseline: текст файлов + список рисков в одном промпте",
        "metrics": {
            "recall": {"value": recall(results), "ci95": bootstrap_ci(results, recall)},
            "precision": {"value": precision(results), "ci95": bootstrap_ci(results, precision)},
            "quotes_found": {"value": _quotes_ok(packages), "ci95": None},
        },
        "errors": [p.error for p in packages if p.error],
    }
    results_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{report['created_at'][:10]}_{set_dir.name}_baseline"
    (results_dir / f"{stem}.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (results_dir / f"{stem}.md").write_text(baseline_markdown(report), encoding="utf-8")
    return report


def _fmt(value: float | None) -> str:
    return "—" if value is None else f"{value:.2f}".replace(".", ",")


def baseline_markdown(report: dict[str, Any]) -> str:
    m = report["metrics"]
    return (
        f"## Базовая линия: набор «{report['dataset']}» "
        f"({plural(report['packages'], ('пакет', 'пакета', 'пакетов'))})\n\n"
        f"Подход: {report['approach']}. Модель: {report['model'] or 'не подключена'}.\n\n"
        "| Метрика | Значение |\n|---|---|\n"
        f"| Полнота | {_fmt(m['recall']['value'])} |\n"
        f"| Точность | {_fmt(m['precision']['value'])} |\n"
        f"| Цитаты найдены в документах | {_fmt(m['quotes_found']['value'])} |\n"
    )
