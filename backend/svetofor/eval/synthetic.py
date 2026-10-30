"""Синтетические пакеты с заложенными рисками.

Шаблоны документов + варианты формулировок («слоты») → документы пакета и `expected.yaml`
с разметкой ВСЕХ правил набора (иначе не посчитать ложные срабатывания) и ожидаемой карточкой.
Формулировки написаны вручную: это эталон, в нём не должно быть случайностей модели.
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any, cast

import yaml

from svetofor.rules.model import Rule
from svetofor.rules.verify import normalize

TEMPLATES = Path(__file__).parent / "templates"
DEFAULTS = {
    "customer": "Государственное казённое учреждение «Ресурсный центр»",
    "place": "634050, г. Томск, пр. Ленина, д. 30",
}


@dataclass
class Trigger:
    rule: str
    quote: str | None


@dataclass
class Block:
    kind: str  # title | subtitle | paragraph | heading | clause | hidden (белым по белому)
    text: str


@dataclass
class GeneratedPackage:
    id: str
    description: str
    documents: dict[str, tuple[str, list[Block]]]  # ключ → (имя файла, блоки)
    triggers: list[Trigger] = field(default_factory=list)
    facts: dict[str, Any] = field(default_factory=dict)


def load_templates() -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    def read(name: str) -> dict[str, Any]:
        data: dict[str, Any] = yaml.safe_load((TEMPLATES / name).read_text(encoding="utf-8"))
        return data

    return read("documents.yaml"), read("slots.yaml"), read("packages.yaml")


@lru_cache(maxsize=1)
def documents_223() -> dict[str, Any]:
    """Шаблоны пакета по 223-ФЗ: запрос котировок, проект договора, техническое задание."""
    data: dict[str, Any] = yaml.safe_load(
        (TEMPLATES / "documents_223.yaml").read_text(encoding="utf-8")
    )
    return data


def money(value: float) -> str:
    """1250000 → «1 250 000,00»."""
    whole, cents = f"{value:.2f}".split(".")
    return f"{int(whole):,}".replace(",", " ") + f",{cents}"


def _params(spec: dict[str, Any], subjects: dict[str, Any], index: int) -> dict[str, Any]:
    params = {**DEFAULTS, **subjects[spec["subject"]], **spec.get("params", {})}
    params.setdefault("number", f"0373100012326000{index:03d}")
    params.setdefault("quantity_spec", params["quantity"])
    params["nmck_text"] = money(params["nmck"])
    params["price_text"] = money(params["nmck"])
    return params


def _fill(text: str, params: dict[str, Any]) -> str:
    return " ".join(str(text).format(**params).split())


def _variant_texts(variant: dict[str, Any], document: str) -> list[str]:
    text = variant.get("text", [])
    if isinstance(text, dict):
        text = text.get(document, [])
    return [text] if isinstance(text, str) else list(text)


def _render(
    key: str, template: dict[str, Any], choose: dict[str, dict[str, Any]], params: dict[str, Any]
) -> list[Block]:
    blocks = [Block("title", template["title"])]
    if template.get("subtitle"):
        blocks.append(Block("subtitle", _fill(template["subtitle"], params)))
    if template.get("preamble"):
        blocks.append(Block("paragraph", _fill(template["preamble"], params)))
    for number, section in enumerate(template["sections"], start=1):
        blocks.append(Block("heading", f"{number}. {section['title']}"))
        clause = 0
        for item in section["items"]:
            variant = choose[item["slot"]] if isinstance(item, dict) else {}
            texts = _variant_texts(variant, key) if isinstance(item, dict) else [item]
            for text in texts:
                if variant.get("hidden"):  # скрытый текст — без номера пункта
                    blocks.append(Block("hidden", _fill(text, params)))
                    continue
                clause += 1
                blocks.append(Block("clause", f"{number}.{clause}. {_fill(text, params)}"))
    return blocks


def build_package(
    spec: dict[str, Any],
    index: int,
    documents: dict[str, Any],
    slots: dict[str, Any],
    subjects: dict[str, Any],
) -> GeneratedPackage:
    params = _params(spec, subjects, index)
    law = spec.get("law", "44-ФЗ")
    if law == "223-ФЗ":  # свои шаблоны: договор по 223-ФЗ не ссылается на нормы 44-ФЗ
        documents = documents_223()
    # только слоты, которые стоят в документах пакета: иначе в ожидаемую карточку попадали бы
    # значения из вариантов, которых в тексте нет
    used = {
        item["slot"]
        for template in documents.values()
        for section in template["sections"]
        for item in section["items"]
        if isinstance(item, dict)
    }
    unknown = set(spec.get("slots", {})) - used
    if unknown:
        raise ValueError(f"{spec['id']}: слоты не стоят в документах пакета: {sorted(unknown)}")
    chosen = {
        name: variants[spec.get("slots", {}).get(name, "safe")]
        for name, variants in slots.items()
        if name in used
    }
    package = GeneratedPackage(spec["id"], spec["description"], {})
    for key, template in documents.items():
        package.documents[key] = (template["filename"], _render(key, template, chosen, params))
    package.facts = {
        "law": law,
        "procurement_number": params["number"],
        "nmck_rub": params["nmck"],
    }
    for variant in chosen.values():
        for trigger in variant.get("triggers", []):
            quote = trigger.get("quote")
            package.triggers.append(
                Trigger(trigger["rule"], _fill(quote, params) if quote else None)
            )
        for fact_key, value in variant.get("facts", {}).items():
            filled = _fill(value, params) if isinstance(value, str) else value
            package.facts[fact_key] = (
                int(filled) if isinstance(filled, str) and filled.isdigit() else filled
            )
    return package


def _document_of(package: GeneratedPackage, quote: str | None) -> str | None:
    if not quote:
        return None
    wanted = normalize(quote)  # как при оценке: латиница-двойник, ё, кавычки не мешают
    for filename, blocks in package.documents.values():
        if any(wanted in normalize(block.text) for block in blocks):
            return filename
    raise ValueError(f"{package.id}: фрагмент «{quote}» не найден ни в одном документе")


def expected(package: GeneratedPackage, rules: list[Rule]) -> dict[str, Any]:
    """Разметка всех правил набора: present true/false, документ и фрагмент цитаты."""
    present = {t.rule: t for t in package.triggers}
    unknown = set(present) - {r.id for r in rules}
    if unknown:
        raise ValueError(f"{package.id}: неизвестные правила {sorted(unknown)}")
    items: list[dict[str, Any]] = []
    for rule in rules:
        trigger = present.get(rule.id)
        if trigger is None:
            items.append({"rule": rule.id, "present": False})
            continue
        item: dict[str, Any] = {"rule": rule.id, "present": True}
        document = _document_of(package, trigger.quote)
        if document:
            item["document"] = document
            item["quote_contains"] = trigger.quote
        items.append(item)
    return {
        "package": package.id,
        "description": package.description,
        "law": "44-ФЗ",
        "annotators": ["генератор (эталон по построению)"],
        "expected": items,
        "facts": package.facts,
    }


# ─────────────────────────── Запись файлов ───────────────────────────


def write_docx(blocks: list[Block], path: Path) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(12)
    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_paragraph()
            paragraph.add_run(block.text).bold = True
        elif block.kind == "hidden":
            run = document.add_paragraph().add_run(block.text)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif block.kind in ("title", "subtitle"):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(block.text)
            run.bold = block.kind == "title"
        else:
            paragraph = document.add_paragraph(block.text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    document.save(str(path))
    return path


def write_pdf(blocks: list[Block], path: Path) -> Path:
    """PDF с переносом строк — для проверки без LibreOffice (нужны шрифты pymupdf-fonts)."""
    import pymupdf

    from svetofor.pipeline.pdf import open_pdf

    font_of = cast(Any, pymupdf).Font
    width, height, margin = 595, 842, 56
    doc = open_pdf()
    page = None
    y: float = height
    for block in blocks:
        bold = block.kind in ("title", "heading")
        size = 14 if block.kind == "title" else 11
        font = "notosbo" if bold else "notos"
        lines = _wrap(block.text, font_of(font), size, width - 2 * margin)
        needed = len(lines) * size * 1.45 + 6
        if page is None or y + needed > height - margin:
            page = doc.new_page(width=width, height=height)
            for name in ("notos", "notosbo"):
                page.insert_font(fontname=name, fontbuffer=font_of(name).buffer)
            y = margin
        for line in lines:
            y += size * 1.45
            x = (
                (width - font_of(font).text_length(line, size)) / 2
                if block.kind in ("title", "subtitle")
                else margin
            )
            color = (1.0, 1.0, 1.0) if block.kind == "hidden" else (0.0, 0.0, 0.0)
            page.insert_text((x, y), line, fontname=font, fontsize=size, color=color)
        y += 6
    doc.save(path)
    doc.close()
    return path


def _wrap(text: str, font: Any, size: float, limit: float) -> list[str]:
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if current and font.text_length(candidate, size) > limit:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def write_package(
    package: GeneratedPackage, rules: list[Rule], folder: Path, fmt: str = "docx"
) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    for old in folder.glob("*"):
        if old.is_file():
            old.unlink()
    writer = write_docx if fmt == "docx" else write_pdf
    for filename, blocks in package.documents.values():
        name = re.sub(r"\.docx$", f".{fmt}", filename)
        writer(blocks, folder / name)
    data = expected(package, rules)
    if fmt != "docx":
        for item in data["expected"]:
            if "document" in item:
                item["document"] = re.sub(r"\.docx$", f".{fmt}", item["document"])
    (folder / "expected.yaml").write_text(
        yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=110), encoding="utf-8"
    )
    return folder


# разметка, спорные случаи и ссылка на закупку в ЕИС — не документы
ANNOTATION_FILES = frozenset({"expected.yaml", "notes.md", "source.txt"})


def package_zip(folder: Path, target: Path) -> Path:
    """Архив пакета для загрузки (как его прислал бы заказчик), без файлов разметки."""
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(folder.iterdir()):
            if path.name not in ANNOTATION_FILES and path.is_file():
                archive.write(path, path.name)
    return target


def generate(out_dir: Path, rules: list[Rule], fmt: str = "docx") -> list[Path]:
    documents, slots, catalog = load_templates()
    folders = []
    for index, spec in enumerate(catalog["packages"], start=1):
        package = build_package(spec, index, documents, slots, catalog["subjects"])
        folders.append(write_package(package, rules, out_dir / spec["id"], fmt))
    return folders


def generate_demo(target: Path, rules: list[Rule], fmt: str = "docx") -> Path:
    """Демо-пакет для samples/: архив документов и рядом — эталон для проверки демо."""
    documents, slots, catalog = load_templates()
    spec = catalog["demo"]
    package = build_package(spec, 777, documents, slots, catalog["subjects"])
    folder = write_package(package, rules, target.parent / spec["id"], fmt)
    package_zip(folder, target)
    return target
