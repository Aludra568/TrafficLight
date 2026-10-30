"""«Изменения в документации» в Word: риски, условия закупки, пункты — было и стало.

Удалённые слова — зачёркнутым красным в колонке «Было», вставленные — жирным зелёным в «Стало»:
так юрист видит правку, не сравнивая два документа глазами.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx
from docx.shared import Cm, Pt, RGBColor

from svetofor.reports.docx_report import (
    DISCLAIMER,
    _shade,
    _table,
    save_atomic,
    setup_document,
    xml_safe,
)

_STATUS = {
    "changed": "Изменён",
    "renumbered": "Перенумерован",
    "added": "Добавлен",
    "removed": "Удалён",
}
_DOC_STATUS = {
    "changed": "есть изменения",
    "same": "без изменений",
    "added": "новый документ",
    "removed": "документ удалён",
}
_SEVERITY = {"red": "Критично", "yellow": "Внимание"}
_RED, _GREEN = RGBColor(0xB0, 0x00, 0x20), RGBColor(0x0B, 0x6E, 0x2B)
_WIDTHS = [2.2, 6.0, 6.0, 2.6]


def _cell_text(cell: Any, text: str, size: float = 9.5) -> None:
    cell.paragraphs[0].add_run(text).font.size = Pt(size)


def _cell_diff(cell: Any, diff: list[dict[str, str]], side: str) -> None:
    """Колонка «Было» (side=delete) или «Стало» (side=insert) с выделением изменённых слов."""
    paragraph = cell.paragraphs[0]
    for part in diff:
        if part["op"] not in ("equal", side):
            continue
        run = paragraph.add_run(part["text"] + " ")
        run.font.size = Pt(9.5)
        if part["op"] == "delete":
            run.font.strike, run.font.color.rgb = True, _RED
        elif part["op"] == "insert":
            run.bold, run.font.color.rgb = True, _GREEN


def _risks(document: Any, risks: dict[str, list[dict[str, Any]]]) -> None:
    document.add_heading("Риски", level=1)
    rows = (
        [("Появился", r) for r in risks["appeared"]]
        + [("Цвет изменился", r) for r in risks["severity_changed"]]
        + [("Ушёл", r) for r in risks["resolved"]]
    )
    if not rows:
        document.add_paragraph("Набор рисков не изменился.")
        return
    widths = [3.0, 3.0, 10.8]
    table = _table(document, ["Что", "Цвет", "Риск"], widths)
    for what, risk in rows:
        cells = table.add_row().cells
        color = _SEVERITY.get(risk["severity"], risk["severity"])
        if "old_severity" in risk:
            color = f"{_SEVERITY.get(risk['old_severity'], '')} → {color}"
        values = [what, color, f"{risk['rule_code']}. {risk['title']}"]
        for cell, value, width in zip(cells, values, widths, strict=True):
            cell.width = Cm(width)
            _cell_text(cell, value)
        if risk["severity"] == "red" and what != "Ушёл":
            _shade(cells[1], "F8D7DA")


def _card(document: Any, card: list[dict[str, Any]]) -> None:
    document.add_heading("Условия закупки", level=1)
    if not card:
        document.add_paragraph("Условия в карточке закупки не изменились.")
        return
    widths = [5.0, 5.9, 5.9]
    table = _table(document, ["Условие", "Было", "Стало"], widths)
    for change in card:
        cells = table.add_row().cells
        for cell, value, width in zip(
            cells,
            [change["label"], change["old"] or "—", change["new"] or "—"],
            widths,
            strict=True,
        ):
            cell.width = Cm(width)
            _cell_text(cell, value)


def _documents(document: Any, documents: list[dict[str, Any]]) -> None:
    document.add_heading("Пункты документов", level=1)
    for item in documents:
        name = item["new_filename"] or item["old_filename"]
        document.add_heading(f"{name} — {_DOC_STATUS[item['status']]}", level=2)
        if not item["changes"]:
            continue
        table = _table(document, ["Пункт", "Было", "Стало", "Изменение"], _WIDTHS)
        for change in item["changes"]:
            cells = table.add_row().cells
            for cell, width in zip(cells, _WIDTHS, strict=True):
                cell.width = Cm(width)
            old, new = change["old"], change["new"]
            number = (new or old)["number"] or "—"
            if change["status"] == "renumbered":
                number = f"{old['number']} → {new['number']}"
            _cell_text(cells[0], number)
            if change["status"] == "changed" and change["diff"]:
                _cell_diff(cells[1], change["diff"], "delete")
                _cell_diff(cells[2], change["diff"], "insert")
            else:
                _cell_text(cells[1], old["text"] if old else "—")
                _cell_text(cells[2], new["text"] if new else "—")
            _cell_text(cells[3], _STATUS[change["status"]])


def render_comparison_docx(diff: dict[str, Any], title: str, path: Path) -> Path:
    diff, title = xml_safe(diff), xml_safe(title)
    document = docx.Document()
    setup_document(document)
    document.add_heading("Изменения в документации закупки", level=0)
    document.add_paragraph(title)
    _risks(document, diff["risks"])
    _card(document, diff["card"])
    _documents(document, diff["documents"])
    document.add_paragraph(
        "Сравнение показывает изменения текста: пункты сопоставлены по номеру, затем по тексту. "
        "Смысл изменений оценивают правила проверки и специалист. " + DISCLAIMER
    )
    return save_atomic(document, path)
