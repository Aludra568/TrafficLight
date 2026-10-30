"""Отчёт в Word.

Титул (закупка, дата, светофор, индекс) → карточка закупки → таблица рисков (цвет, правило,
цитата, место, норма, рекомендация) → что проверено → оговорка. Оформление деловое: юрист
пересылает отчёт руководителю без правок. Цвет всегда продублирован словом.
"""

from __future__ import annotations

import os
import re
import tempfile
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, cast

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel

from svetofor.db.enums import FindingKind, Light, RuleOutcome, RuleStatus, Severity
from svetofor.schemas.checks import FindingOut, ReportOut

FONT = "Times New Roman"
DISCLAIMER = (
    "Отчёт — подсказка для специалиста, а не юридическое заключение. Цитаты взяты из документов "
    "и проверены программно; правила со статусом «черновик» ещё не проверены юристом."
)
LIGHT_WORDS = {
    Light.RED: "Красный — есть критичные риски",
    Light.YELLOW: "Жёлтый — есть риски, требующие внимания",
    Light.GREEN: "Зелёный — риски не найдены",
    Light.GRAY: "Серый — риски не оценены: проверена меньшая часть правил",
}
SEVERITY_WORDS = {Severity.RED: "Критично", Severity.YELLOW: "Внимание"}
SEVERITY_FILL = {Severity.RED: "F8D7DA", Severity.YELLOW: "FFF3CD"}
OUTCOME_WORDS = {
    RuleOutcome.VIOLATION: "Риск найден",
    RuleOutcome.ABSENT: "Условие не найдено",
    RuleOutcome.OK: "Риска нет",
    RuleOutcome.NOT_CHECKED: "Не проверено",
}
KIND_WORDS = {
    FindingKind.CONTRADICTION: "Противоречие между документами",
    FindingKind.ABSENT: "Условие не найдено",
}


# Word (XML 1.0) не допускает управляющих символов: python-docx падает на первом же — а они
# приходят в именах файлов из архивов, в текстах правил методолога и в ответах модели
_XML_INVALID = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


def xml_safe[T](value: T) -> T:
    """Те же данные без символов, которых не бывает в XML: строки, списки, словари, модели."""
    if isinstance(value, str):
        return cast(T, _XML_INVALID.sub("", value))
    if isinstance(value, BaseModel):
        return cast(T, type(value).model_validate(xml_safe(value.model_dump())))
    if isinstance(value, list | tuple):
        return cast(T, type(value)(xml_safe(item) for item in value))
    if isinstance(value, dict):
        return cast(T, {key: xml_safe(item) for key, item in value.items()})
    return value


def save_atomic(document: Any, path: Path) -> Path:
    """Сохранить через временный файл: одновременная выгрузка того же отчёта (второй щелчок,
    PDF из этого же Word) не отдаст и не прочитает наполовину записанный файл."""
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, name = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
    os.close(handle)
    temporary = Path(name)
    try:
        document.save(str(temporary))
        try:
            temporary.replace(path)
        except PermissionError:
            # Windows (режим разработки): прежний файл ещё отдаётся другому запросу и не
            # заменяется — этот запрос получит свою копию (в Linux замена открытого файла можно)
            path = path.with_stem(f"{path.stem}-{uuid.uuid4().hex[:8]}")
            temporary.replace(path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return path


def setup_document(document: Any) -> None:
    section = document.sections[0]
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(2)
    style = document.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Title"):
        heading = document.styles[name]
        heading.font.name, heading.font.color.rgb = FONT, RGBColor(0, 0, 0)
        heading.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _shade(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _table(document: Any, header: list[str], widths: list[float]) -> Any:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text, width in zip(table.rows[0].cells, header, widths, strict=True):
        cell.width = Cm(width)
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        _shade(cell, "E7E6E6")
    return table


def _row(table: Any, values: list[str], widths: list[float], size: float = 9.5) -> Any:
    cells = table.add_row().cells
    for cell, value, width in zip(cells, values, widths, strict=True):
        cell.width = Cm(width)
        cell.paragraphs[0].add_run(value).font.size = Pt(size)
    return cells


def _place(finding: FindingOut) -> str:
    if finding.evidence:
        return "; ".join(f"{e.section_path} (стр. {e.page_from})" for e in finding.evidence)
    if finding.search_scope:
        return f"Проверено: {finding.search_scope}"
    if finding.checked_documents:
        return "Искали в: " + ", ".join(finding.checked_documents)
    return finding.section_path or "—"


def _norm(finding: FindingOut) -> str:
    refs = [" ".join(x for x in (ref.act, ref.ref) if x) for ref in finding.legal_basis]
    unverified = finding.rule_status is not RuleStatus.VERIFIED
    return ("; ".join(refs) or "—") + ("\n(правило не проверено юристом)" if unverified else "")


def _title(document: Any, report: ReportOut) -> None:
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Отчёт о проверке закупочной документации")
    run.bold, run.font.size = True, Pt(15)
    lines = [
        ("Пакет", report.package.title),
        ("Файл", report.package.source_filename),
        (
            "Дата проверки",
            (report.check.finished_at or report.check.created_at).strftime("%d.%m.%Y %H:%M"),
        ),
        ("Итог", LIGHT_WORDS.get(report.check.light, "—") if report.check.light else "—"),
        (
            "Индекс риска",
            f"{report.risk.index:g} из 100 — {report.risk.recommendation_text}"
            if report.risk.index is not None
            else "—",
        ),
        ("Набор правил", report.check.ruleset or "—"),
    ]
    for label, value in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)


def _card(document: Any, report: ReportOut) -> None:
    document.add_heading("Карточка закупки", level=1)
    if not report.card:
        document.add_paragraph("Ключевые условия в документах не найдены.")
        return
    widths = [4.5, 4.0, 7.5]
    table = _table(document, ["Условие", "Значение", "Где сказано"], widths)
    for item in report.card:
        _row(
            table,
            [item.label, item.display, f"«{item.quote}» ({item.filename}, стр. {item.page})"],
            widths,
        )


def _risks(document: Any, report: ReportOut) -> None:
    document.add_heading("Риски", level=1)
    if not report.findings:
        document.add_paragraph("Риски по правилам набора не найдены.")
        return
    widths = [2.0, 3.6, 4.6, 2.8, 3.0]
    table = _table(document, ["Цвет", "Риск", "Цитата и место", "Норма", "Что делать"], widths)
    for finding in report.findings:
        risk = f"{finding.rule_code}. {finding.title}\n{finding.message}"
        if finding.kind in KIND_WORDS:
            risk = f"{KIND_WORDS[finding.kind]}. {risk}"
        if finding.needs_review:
            comment = finding.review.comment if finding.review else None
            risk += "\nТребует проверки специалистом: второе мнение расходится" + (
                f" ({comment})" if comment else ""
            )
        quotes = "\n".join(f"«{e.quote}»" for e in finding.evidence) or "—"
        cells = _row(
            table,
            [
                SEVERITY_WORDS[finding.severity],
                risk,
                f"{quotes}\n{_place(finding)}",
                _norm(finding),
                finding.recommendation,
            ],
            widths,
        )
        _shade(cells[0], SEVERITY_FILL[finding.severity])
        cells[0].paragraphs[0].runs[0].bold = True


def _checked(document: Any, report: ReportOut) -> None:
    document.add_heading("Что проверено", level=1)
    summary = report.rules_summary
    document.add_paragraph(
        f"Правил: {summary.checked}. Риск найден: {summary.violations}, условие не найдено: "
        f"{summary.absent}, риска нет: {summary.ok}, не проверено: {summary.not_checked}."
    )
    rejected = report.counters.get("rejected_unverified", 0)
    if rejected:
        document.add_paragraph(
            f"Отклонено как неподтверждённые цитатой: {rejected} (в отчёт не включены)."
        )
    widths = [2.0, 8.0, 3.0, 3.0]
    table = _table(document, ["Код", "Правило", "Результат", "Причина"], widths)
    for rule in report.rules:
        _row(
            table,
            [rule.rule_code, rule.title, OUTCOME_WORDS[rule.outcome], rule.reason or ""],
            widths,
        )


def render_report_docx(report: ReportOut, path: Path) -> Path:
    report = xml_safe(report)
    document = docx.Document()
    setup_document(document)
    _title(document, report)
    _card(document, report)
    _risks(document, report)
    _checked(document, report)
    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(12)
    run = disclaimer.add_run(DISCLAIMER)
    run.italic, run.font.size = True, Pt(9)
    document.core_properties.title = f"Проверка: {report.package.title}"
    document.core_properties.created = datetime.now()
    return save_atomic(document, path)
