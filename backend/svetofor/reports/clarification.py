"""Черновик запроса на разъяснение положений документации — из найденных рисков.

Без модели: вопросы собираются из замечаний (точные цитаты, место) и формулировки правила
`clarification`, которую ведёт методолог. Если её нет — вопрос строится по виду замечания.
Документ — черновик: реквизиты заполняет и проверяет специалист.
"""

from __future__ import annotations

import re
import uuid
from collections.abc import Collection
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from svetofor.db.enums import FindingKind, Severity
from svetofor.reports.docx_report import save_atomic, setup_document, xml_safe
from svetofor.rules.model import Rule
from svetofor.rules.templates import render
from svetofor.schemas.checks import FindingOut, ReportOut

DRAFT_NOTE = (
    "Черновик подготовлен системой «Светофор» по результатам проверки. Проверьте формулировки, "
    "заполните реквизиты и удалите эту строку перед отправкой."
)
_GENERIC = {
    FindingKind.CONTRADICTION: "Просим указать, какое из значений является верным, "
    "и внести соответствующие изменения в документацию.",
    FindingKind.ABSENT: "Просим разъяснить, предусмотрено ли это условие документацией, "
    "и указать, в каком разделе.",
    FindingKind.VIOLATION: "Просим разъяснить данное положение и подтвердить его соответствие "
    "требованиям законодательства о контрактной системе.",
}
_CLAUSE_NUMBER = re.compile(r"\s*(\d{1,2}(?:\.\d{1,3}){0,4}\.?)\s+(?=[А-ЯЁA-Z«])")


_DATE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")
# ч. 5 ст. 42 44-ФЗ: не позднее чем за три дня до окончания подачи заявок, не более трёх запросов
REQUEST_DAYS_BEFORE = 3
REQUEST_LIMIT_NOTE = (
    "По одной закупке можно направить не более трёх запросов — вопросы собраны в один."
)


def _under_44fz(report: ReportOut) -> bool:
    """Сроки и лимит запросов — нормы 44-ФЗ; закон не определён — считаем 44-ФЗ (основной)."""
    return _card_value(report, "law") in (None, "44-ФЗ")


def request_deadline(report: ReportOut) -> str | None:
    """Последний день запроса о разъяснениях по 44-ФЗ — из даты окончания подачи заявок."""
    if not _under_44fz(report):
        return None
    match = _DATE.search(_card_value(report, "application_deadline") or "")
    if match is None:
        return None
    try:
        end = date(int(match[3]), int(match[2]), int(match[1]))
    except ValueError:  # «31.02.2026» после распознавания — даты нет
        return None
    return f"{end - timedelta(days=REQUEST_DAYS_BEFORE):%d.%m.%Y}"


def deadline_note(report: ReportOut) -> str | None:
    """Подсказка о сроке и числе запросов — только для 44-ФЗ: в 223-ФЗ и коммерческих
    закупках порядок задаёт положение о закупке, общий закон о трёх запросах не действует."""
    if not _under_44fz(report):
        return None
    deadline = request_deadline(report)
    if deadline is None:
        return REQUEST_LIMIT_NOTE
    return (
        f"Направьте запрос через электронную площадку не позднее {deadline} — за три дня до "
        f"окончания подачи заявок (ч. 5 ст. 42 44-ФЗ). {REQUEST_LIMIT_NOTE}"
    )


class NothingToAskError(ValueError):
    """Нет замечаний, по которым можно задать вопрос."""


def question(finding: FindingOut, rule: Rule | None) -> str:
    if rule is not None and rule.clarification:
        return render(rule.clarification, quote=finding.quote or "", value=finding.value or "")
    return f"{finding.message} {_GENERIC[finding.kind]}"


def quoted(text: str) -> str:
    """Цитата в кавычках по правилам русской пунктуации: точка внутри кавычек не ставится."""
    return "«" + " ".join(text.split()).rstrip(".;,") + "»"


def without_clause_number(quote: str, section_path: str) -> str:
    """«3.1. Срок поставки…» → «Срок поставки…», если пункт 3.1 уже назван в месте цитаты."""
    match = _CLAUSE_NUMBER.match(quote)
    if match and section_path.rsplit("›", 1)[-1].strip() == match.group(1).rstrip("."):
        return quote[match.end() :]
    return quote


def _where(finding: FindingOut, filenames: dict[uuid.UUID, str]) -> str:
    parts = [
        f"в документе «{filenames.get(e.document_id, 'документ')}» ({e.section_path}, "
        f"стр. {e.page_from}) указано: {quoted(without_clause_number(e.quote, e.section_path))}"
        for e in finding.evidence
    ]
    if not parts:
        return ""
    text = ", а ".join(parts)
    return text[:1].upper() + text[1:] + ". "  # не .capitalize(): он опустил бы регистр цитаты


def worth_asking(finding: FindingOut, rule: Rule | None) -> bool:
    """По умолчанию — только то, что действительно спрашивают у заказчика.

    Противоречие и ненайденное условие — всегда вопрос; критичный риск — вопрос о законности;
    правило с формулировкой вопроса — решение методолога. «Аванс не предусмотрен» — коммерческий
    риск, а не повод для запроса: такие замечания специалист добавляет вручную.
    """
    if finding.needs_review:  # второе мнение расходится — сначала решает специалист
        return False
    if finding.kind in (FindingKind.CONTRADICTION, FindingKind.ABSENT):
        return True
    return finding.severity is Severity.RED or bool(rule and rule.clarification)


def select_findings(
    report: ReportOut, rules: dict[str, Rule], finding_ids: Collection[uuid.UUID] | None
) -> list[FindingOut]:
    if finding_ids is None:
        chosen = [f for f in report.findings if worth_asking(f, rules.get(f.rule_code))]
    else:
        chosen = [f for f in report.findings if f.id in finding_ids]
    if not chosen:
        raise NothingToAskError("Нет замечаний для запроса на разъяснение")
    return chosen


def _card_value(report: ReportOut, key: str) -> str | None:
    return next((str(item.value) for item in report.card if item.key == key), None)


def _paragraph(document: Any, text: str, *, bold: bool = False, align: Any = None) -> Any:
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(xml_safe(text))  # шаблоны вопросов пишет методолог
    run.bold = bold
    return paragraph


def render_clarification_docx(
    report: ReportOut,
    rules: dict[str, Rule],
    path: Path,
    finding_ids: Collection[uuid.UUID] | None = None,
) -> Path:
    findings = select_findings(report, rules, finding_ids)
    filenames = {d.id: d.filename for d in report.documents}
    document = docx.Document()
    setup_document(document)
    for text in (DRAFT_NOTE, deadline_note(report)):
        if text is None:
            continue
        note = document.add_paragraph().add_run(xml_safe(text))
        note.italic, note.font.size, note.font.color.rgb = True, Pt(9), RGBColor(0xC0, 0, 0)
    for line in (
        "Заказчику: ______________________________",
        "От участника закупки: ______________",
    ):
        _paragraph(document, line, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _paragraph(document, "ЗАПРОС", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _paragraph(
        document,
        "о даче разъяснений положений документации о закупке",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    number = _card_value(report, "procurement_number") or "____________"
    subject = _card_value(report, "subject")
    intro = f"Закупка № {number}" + (f", предмет закупки: {subject}" if subject else "") + "."
    _paragraph(document, intro)
    _paragraph(document, "Просим дать разъяснения по следующим положениям документации:")
    for index, finding in enumerate(findings, start=1):
        rule = rules.get(finding.rule_code)
        paragraph = _paragraph(
            document, f"{index}. {_where(finding, filenames)}{question(finding, rule)}"
        )
        paragraph.paragraph_format.first_line_indent = Cm(1)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _paragraph(
        document,
        "Ответ просим направить в порядке и сроки, установленные законодательством "
        "о контрактной системе.",
    )
    for line in (
        "",
        "______________________ / ______________________ /",
        "(должность, подпись, расшифровка)",
        "«____» ______________ 20___ г.",
    ):
        _paragraph(document, line)
    document.core_properties.title = "Запрос о даче разъяснений положений документации"
    return save_atomic(document, path)
