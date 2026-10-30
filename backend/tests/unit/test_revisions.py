"""Неподтверждённые исправления Word принимаются до конвертации: удалённый текст не читается."""

from __future__ import annotations

import zipfile
from pathlib import Path

import docx

from svetofor.pipeline.revisions import accept_docx_revisions, accept_revisions_xml

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
REVISED_BODY = (
    f'<w:p xmlns:w="{W}"><w:r><w:t xml:space="preserve">Оплата в течение </w:t></w:r>'
    '<w:del w:id="1" w:author="Заказчик"><w:r><w:delText>30</w:delText></w:r></w:del>'
    '<w:ins w:id="2" w:author="Заказчик"><w:r><w:t>15</w:t></w:r></w:ins>'
    '<w:r><w:t xml:space="preserve"> рабочих дней.</w:t></w:r></w:p>'
)


def _docx_with_revisions(path: Path) -> Path:
    document = docx.Document()
    document.add_paragraph("ПРОЕКТ КОНТРАКТА")
    document.add_paragraph("МАРКЕР")
    document.save(str(path))
    with zipfile.ZipFile(path) as archive:
        parts = {i.filename: archive.read(i.filename) for i in archive.infolist()}
    xml = parts["word/document.xml"].decode("utf-8")
    body = REVISED_BODY.replace(f' xmlns:w="{W}"', "")
    start = xml.index("<w:p>", xml.index("ПРОЕКТ КОНТРАКТА"))
    end = xml.index("</w:p>", xml.index("МАРКЕР")) + len("</w:p>")
    parts["word/document.xml"] = (xml[:start] + body + xml[end:]).encode("utf-8")
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return path


def test_xml_deleted_text_gone_inserted_kept() -> None:
    xml, count = accept_revisions_xml(REVISED_BODY.encode("utf-8"))
    text = xml.decode("utf-8")
    assert count == 2 and "30" not in text and ">15<" in text
    assert "w:ins" not in text and "w:del" not in text


def test_xml_without_revisions_untouched() -> None:
    original = f'<w:p xmlns:w="{W}"><w:r><w:t>Без исправлений</w:t></w:r></w:p>'.encode()
    assert accept_revisions_xml(original) == (original, 0)


def test_docx_copy_reads_as_accepted(tmp_path: Path) -> None:
    src = _docx_with_revisions(tmp_path / "contract.docx")
    dest = tmp_path / "accepted.docx"
    assert accept_docx_revisions(src, dest) == 2
    text = "\n".join(p.text for p in docx.Document(str(dest)).paragraphs)
    assert "Оплата в течение 15 рабочих дней." in text and "30" not in text


def test_clean_docx_not_copied(tmp_path: Path) -> None:
    src = tmp_path / "clean.docx"
    document = docx.Document()
    document.add_paragraph("Обычный текст")
    document.save(str(src))
    dest = tmp_path / "accepted.docx"
    assert accept_docx_revisions(src, dest) == 0 and not dest.exists()
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"not a zip")
    assert accept_docx_revisions(broken, dest) == 0
