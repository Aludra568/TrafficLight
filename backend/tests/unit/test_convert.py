"""Конвертация в PDF: проверки PDF, картинки, LibreOffice (в Docker), заголовки из стилей DOCX."""

from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from svetofor.pipeline.convert import (
    IMAGE_EXTENSIONS,
    OFFICE_EXTENSIONS,
    PDF_EXTENSIONS,
    XML_EXTENSIONS,
    ConvertError,
    _decode_text,
    convert_to_pdf,
    docx_heading_hints,
    soffice_command,
    xml_text,
)
from svetofor.pipeline.extract import extract_document
from svetofor.pipeline.unpack import SUPPORTED_EXTENSIONS
from tests.support.pdfgen import L, PdfSpec, build_pdf


def test_every_supported_extension_has_a_converter() -> None:
    assert (
        PDF_EXTENSIONS | OFFICE_EXTENSIONS | IMAGE_EXTENSIONS | XML_EXTENSIONS
        == SUPPORTED_EXTENSIONS
    )


# ─── XML из ЕИС ───


def test_xml_values_become_text_lines() -> None:
    raw = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<ns2:export xmlns:ns2="http://zakupki.gov.ru/oos/export/1">'
        "<purchaseObjectInfo>Поставка   ноутбуков</purchaseObjectInfo><!-- служебное -->"
        "<maxPrice>1890000.00</maxPrice><term>15 рабочих дней</term></ns2:export>"
    ).encode()
    assert xml_text(raw) == "Поставка ноутбуков\n1890000.00\n15 рабочих дней"


@pytest.mark.parametrize("encoding", ["utf-16-le", "utf-16-be"])
@pytest.mark.parametrize("declaration", ["", '<?xml version="1.0" encoding="UTF-16"?>'])
def test_xml_in_utf16_without_bom(encoding: str, declaration: str) -> None:
    """ЕИС отдаёт «Электронный документ, полученный из внешней системы.xml» в UTF-16 без метки
    порядка байт — libxml2 считал такой файл повреждённым (закупка 0120100008026000013)."""
    xml = f"{declaration}<purchaseresults><purchname>поставка картриджей</purchname>"
    xml += "<economypercent>10.00</economypercent></purchaseresults>"
    assert xml_text(xml.encode(encoding)) == "поставка картриджей\n10.00"


def test_xml_external_entities_are_not_read(tmp_path: Path) -> None:
    """Чужой XML не читает файлы сервера и не ходит в сеть (XXE)."""
    secret = tmp_path / "secret.txt"
    secret.write_text("СЕКРЕТ-СЕРВЕРА", encoding="utf-8")
    raw = (
        f'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "{secret.as_uri()}">'
        '<!ENTITY n SYSTEM "http://127.0.0.1:9/leak">]><d>до &x; &n; после</d>'
    ).encode()
    text = xml_text(raw)
    assert "СЕКРЕТ" not in text and "до" in text and "после" in text


def test_xml_entity_bomb_is_rejected() -> None:
    """«Миллиард смешков»: libxml2 отказывается раздувать сущности — файл отклонён с причиной,
    а не 10 миллионов «lol» в памяти воркера."""
    lol = "".join(f'<!ENTITY l{i} "{f"&l{i - 1};" * 10}">' for i in range(1, 8))
    raw = f'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY l0 "lol">{lol}]><d>&l7;</d>'.encode()
    with pytest.raises(ConvertError, match="XML-файл повреждён"):
        xml_text(raw)


def test_broken_xml_has_a_clear_reason() -> None:
    with pytest.raises(ConvertError, match="XML-файл повреждён"):
        xml_text(b"<d><unclosed></d>")


def test_xml_goes_through_text_conversion(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    seen: dict[str, str] = {}

    def fake_office(src: Path, ext: str, dest: Path, *_: object) -> None:
        seen["ext"], seen["text"] = ext, src.read_text(encoding="utf-8")
        build_pdf(PdfSpec(pages=[[L("текст")]]), dest)

    from svetofor.pipeline import convert

    monkeypatch.setattr(convert, "_office_to_pdf", fake_office)
    src = tmp_path / "Электронный документ.xml"
    src.write_bytes("<d><a>Срок оплаты 30 дней</a></d>".encode())
    assert convert_to_pdf(src, ".xml", tmp_path / "out.pdf") == 1
    assert seen == {"ext": ".txt", "text": "Срок оплаты 30 дней"}
    empty = tmp_path / "empty.xml"
    empty.write_bytes(b"<d><a/></d>")
    with pytest.raises(ConvertError, match="нет текста"):
        convert_to_pdf(empty, ".xml", tmp_path / "out2.pdf")


# ─── PDF ───


def text_pdf(tmp_path: Path, name: str = "src.pdf") -> Path:
    return build_pdf(PdfSpec(pages=[[L("Проект контракта")], [L("Страница два")]]), tmp_path / name)


def test_pdf_copied(tmp_path: Path) -> None:
    dest = tmp_path / "out" / "doc.pdf"
    assert convert_to_pdf(text_pdf(tmp_path), ".PDF", dest) == 2
    assert "Проект контракта" in extract_document(dest).text


def test_pdf_with_permissions_password_is_decrypted(tmp_path: Path) -> None:
    src = text_pdf(tmp_path)
    locked = tmp_path / "locked.pdf"
    with pymupdf.open(src) as doc:
        doc.save(
            locked,
            encryption=pymupdf.PDF_ENCRYPT_AES_256,
            owner_pw="owner-secret",
            permissions=pymupdf.PDF_PERM_ACCESSIBILITY,
        )
    dest = tmp_path / "doc.pdf"
    assert convert_to_pdf(locked, ".pdf", dest) == 2
    with pymupdf.open(dest) as doc:
        assert not doc.is_encrypted


def test_pdf_with_open_password_rejected(tmp_path: Path) -> None:
    locked = tmp_path / "locked.pdf"
    with pymupdf.open(text_pdf(tmp_path)) as doc:
        doc.save(locked, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
    with pytest.raises(ConvertError, match="защищён паролем"):
        convert_to_pdf(locked, ".pdf", tmp_path / "doc.pdf")


def test_broken_pdf(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4 not really")
    with pytest.raises(ConvertError, match="повреждён"):
        convert_to_pdf(broken, ".pdf", tmp_path / "doc.pdf")


# ─── Картинки ───


@pytest.mark.parametrize(
    ("ext", "mode", "fmt"),
    [
        (".png", "RGB", "PNG"),
        (".png", "RGBA", "PNG"),
        (".jpg", "RGB", "JPEG"),
        (".bmp", "RGB", "BMP"),
    ],
)
def test_image_to_pdf(tmp_path: Path, ext: str, mode: str, fmt: str) -> None:
    src = tmp_path / f"scan{ext}"
    Image.new(mode, (200, 300), "white").save(src, format=fmt)
    assert convert_to_pdf(src, ext, tmp_path / "doc.pdf") == 1


def test_multipage_tiff(tmp_path: Path) -> None:
    src = tmp_path / "scan.tiff"
    frames = [Image.new("RGB", (100, 140), color) for color in ("white", "gray")]
    frames[0].save(src, format="TIFF", save_all=True, append_images=frames[1:])
    assert convert_to_pdf(src, ".tiff", tmp_path / "doc.pdf") == 2


def test_broken_image(tmp_path: Path) -> None:
    src = tmp_path / "scan.png"
    src.write_bytes(b"not an image")
    with pytest.raises(ConvertError, match="изображение"):
        convert_to_pdf(src, ".png", tmp_path / "doc.pdf")


def test_unsupported(tmp_path: Path) -> None:
    with pytest.raises(ConvertError, match="не поддерживается"):
        convert_to_pdf(tmp_path / "x.exe", ".exe", tmp_path / "doc.pdf")


# ─── LibreOffice ───


def test_soffice_missing_gives_clear_error(tmp_path: Path) -> None:
    src = tmp_path / "a.docx"
    docx.Document().save(str(src))
    with pytest.raises(ConvertError, match="не установлен LibreOffice"):
        convert_to_pdf(src, ".docx", tmp_path / "doc.pdf", soffice="svetofor-no-such-soffice")


def test_soffice_command(tmp_path: Path) -> None:
    command = soffice_command("soffice", tmp_path / "input.txt", tmp_path / "out", tmp_path / "p")
    assert command[0] == "soffice" and command[-1].endswith("input.txt")
    assert command[1].startswith("-env:UserInstallation=file:")
    assert "--headless" in command and "--infilter=Text (encoded):UTF8" in command
    docx_command = soffice_command("soffice", tmp_path / "input.docx", tmp_path, tmp_path)
    assert not any(part.startswith("--infilter") for part in docx_command)


def test_spreadsheet_exported_one_page_per_sheet(tmp_path: Path) -> None:
    """Широкая спецификация XLSX не режется по столбцам на разные страницы."""
    command = soffice_command("soffice", tmp_path / "input.xlsx", tmp_path, tmp_path)
    target = command[command.index("--convert-to") + 1]
    assert target.startswith("pdf:calc_pdf_Export:") and "SinglePageSheets" in target
    fallback = soffice_command("soffice", tmp_path / "input.xlsx", tmp_path, tmp_path, plain=True)
    assert fallback[fallback.index("--convert-to") + 1] == "pdf"
    docx_command = soffice_command("soffice", tmp_path / "input.docx", tmp_path, tmp_path)
    assert docx_command[docx_command.index("--convert-to") + 1] == "pdf"


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "cp1251"])
def test_decode_text(encoding: str) -> None:
    assert _decode_text("Срок поставки — 10 дней".encode(encoding)) == "Срок поставки — 10 дней"


@pytest.mark.docker("soffice")
def test_docx_converted_by_libreoffice(tmp_path: Path) -> None:
    src = tmp_path / "Проект контракта.docx"
    document = docx.Document()
    document.add_heading("1. Предмет контракта", level=1)
    document.add_paragraph("1.1. Поставщик обязуется поставить товар.")
    document.save(str(src))
    dest = tmp_path / "doc.pdf"
    assert convert_to_pdf(src, ".docx", dest, work_dir=tmp_path) >= 1
    assert "Поставщик обязуется поставить товар" in extract_document(dest).text


@pytest.mark.docker("soffice")
def test_txt_cp1251_converted(tmp_path: Path) -> None:
    src = tmp_path / "письмо.txt"
    src.write_bytes("Заказчик просит уточнить срок поставки.".encode("cp1251"))
    dest = tmp_path / "doc.pdf"
    convert_to_pdf(src, ".txt", dest, work_dir=tmp_path)
    assert "Заказчик просит уточнить срок поставки" in extract_document(dest).text


# ─── Заголовки из стилей DOCX ───


def test_docx_heading_hints(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("Общие положения", level=1)
    document.add_paragraph("Обычный текст пункта.")
    document.add_heading("Требования  к   упаковке", level=2)
    russian = document.styles.add_style("Заголовок 3", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Маркировка", style=russian)
    outlined = document.add_paragraph("Порядок приёмки")
    ppr = outlined._p.get_or_add_pPr()
    level = OxmlElement("w:outlineLvl")
    level.set(qn("w:val"), "0")
    ppr.append(level)
    path = tmp_path / "spec.docx"
    document.save(str(path))
    assert docx_heading_hints(path) == {
        "Общие положения": 1,
        "Требования к упаковке": 2,
        "Маркировка": 3,
        "Порядок приёмки": 1,
    }


def test_docx_heading_hints_on_broken_file(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK not a docx")
    assert docx_heading_hints(path) == {}


def test_office_conversion_runs_without_app_secrets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Разведка 4: LibreOffice открывает чужой документ — секретов в его окружении быть не
    должно (CVE-2024-12426: документ отправлял переменные окружения наружу), профиль — в
    «скрытном режиме» до первого запуска."""
    import subprocess
    from urllib.parse import urlparse
    from urllib.request import url2pathname

    from svetofor.pipeline import convert

    monkeypatch.setenv("LLM_API_KEY", "sk-secret")
    monkeypatch.setenv("DATABASE_URL", "postgresql://u:pass@db/x")
    monkeypatch.setenv("JWT_SECRET", "jwt-secret")
    seen: dict[str, object] = {}

    def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess[bytes]:
        seen["env"] = kwargs["env"]
        profile = Path(url2pathname(urlparse(command[1].split("=", 1)[1]).path))
        seen["profile"] = (profile / "user" / "registrymodifications.xcu").read_text(
            encoding="utf-8"
        )
        out = Path(command[command.index("--outdir") + 1])
        out.mkdir(parents=True, exist_ok=True)
        (out / "input.pdf").write_bytes(b"%PDF-1.4 converted")
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr(convert.subprocess, "run", fake_run)
    source = tmp_path / "Проект контракта.docx"
    source.write_bytes(b"PK fake docx")
    convert.office_to_pdf(source, tmp_path / "out.pdf", work_dir=tmp_path)
    env = seen["env"]
    assert isinstance(env, dict)
    assert not {"LLM_API_KEY", "DATABASE_URL", "JWT_SECRET"} & set(env)
    assert "PATH" in env
    assert "BlockUntrustedRefererLinks" in str(seen["profile"])
    assert "DisableMacrosExecution" in str(seen["profile"])
