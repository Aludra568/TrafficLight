"""Сохранение выгрузок Word: атомарно, а на Windows — и когда прежний файл ещё открыт."""

from __future__ import annotations

from pathlib import Path

import docx
import pytest

from svetofor.reports.docx_report import save_atomic


def test_saved_atomically_without_leftovers(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Отчёт")
    target = tmp_path / "reports" / "r.docx"
    assert save_atomic(document, target) == target
    assert docx.Document(str(target)).paragraphs[0].text == "Отчёт"
    assert [p.name for p in target.parent.iterdir()] == ["r.docx"]


def test_locked_target_on_windows_gets_own_copy(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Код-ревью 2026-10-08: на Windows замена файла, который отдаёт другой запрос, падала 500."""
    target = tmp_path / "r.docx"
    target.write_bytes(b"old")
    real_replace = Path.replace

    def locked(self: Path, other: Path) -> Path:
        if Path(other) == target:
            raise PermissionError("файл открыт другим процессом")
        return real_replace(self, other)

    monkeypatch.setattr(Path, "replace", locked)
    saved = save_atomic(docx.Document(), target)
    assert saved != target and saved.parent == tmp_path and saved.suffix == ".docx"
    assert target.read_bytes() == b"old" and saved.is_file()
