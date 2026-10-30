"""Выгрузка отчёта (Word, PDF) и черновик запроса на разъяснение."""

from __future__ import annotations

import io
import shutil
import time
from typing import Any

import docx
import pytest

from tests.integration.conftest import auth
from tests.integration.test_api_checks import upload

pytestmark = pytest.mark.pg


def docx_text(content: bytes) -> str:
    document = docx.Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_word_report(api: dict[str, Any]) -> None:
    created = upload(api)
    response = api["client"].post(f"/api/checks/{created['check_id']}/export", headers=auth(api))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in response.headers["content-disposition"]
    text = docx_text(response.content)
    for expected in (
        "Отчёт о проверке закупочной документации",
        "Карточка закупки",
        "15 раб. дн.",
        "Критично",
        "Что проверено",
        "не юридическое заключение",
        "правило не проверено юристом",
        "в течение 15 рабочих дней",
    ):
        assert expected in text, expected
    report = (
        api["client"].get(f"/api/checks/{created['check_id']}/report", headers=auth(api)).json()
    )
    table = docx.Document(io.BytesIO(response.content)).tables[1]
    assert len(table.rows) == len(report["findings"]) + 1  # шапка + по строке на риск


def test_pdf_reports_clear_failure_without_libreoffice(api: dict[str, Any]) -> None:
    if shutil.which("soffice"):
        pytest.skip("LibreOffice есть — проверяется успешный путь")
    created = upload(api)
    url = f"/api/checks/{created['check_id']}/export"
    assert api["client"].post(f"{url}?format=pdf", headers=auth(api)).status_code == 202
    ready = api["client"].get(f"{url}/pdf", headers=auth(api))
    assert ready.status_code == 409
    assert ready.json()["error"]["message"].startswith("Не удалось подготовить PDF")


@pytest.mark.docker("soffice")
def test_pdf_report(api: dict[str, Any]) -> None:
    created = upload(api)
    url = f"/api/checks/{created['check_id']}/export"
    api["client"].post(f"{url}?format=pdf", headers=auth(api))
    for _ in range(100):
        ready = api["client"].get(f"{url}/pdf", headers=auth(api))
        if ready.status_code != 202:
            break
        time.sleep(0.2)
    assert ready.status_code == 200 and ready.content.startswith(b"%PDF")


def test_clarification_draft(api: dict[str, Any]) -> None:
    created = upload(api)
    client = api["client"]
    url = f"/api/checks/{created['check_id']}/clarification"
    response = client.post(url, headers=auth(api), json={})
    assert response.status_code == 200
    text = docx_text(response.content)
    assert "ЗАПРОС" in text and "о даче разъяснений положений документации" in text
    assert "Черновик подготовлен системой" in text
    # по умолчанию: противоречие, не найденное условие, критичный риск — но не «аванса нет»
    # цитата — предложение целиком (канал грамматик), регистр сохранён, номер пункта не повторён
    assert "указано: «Срок поставки: 10 календарных дней с даты заключения контракта»" in text
    assert "«3.1." not in text and "«1.1." not in text
    assert "(Проект контракта › 3. СРОКИ ПОСТАВКИ › 3.1, стр. 1)" in text  # место до пункта
    assert "Ответственность заказчика" in text and "пеню в размере 1/100" in text
    assert "Аванс не предусмотрен" not in text
    findings = client.get(f"/api/checks/{created['check_id']}/report", headers=auth(api)).json()[
        "findings"
    ]
    volume = next(f for f in findings if f["rule_code"] == "VOL-01")
    one = client.post(url, headers=auth(api), json={"finding_ids": [volume["id"]]})
    one_text = docx_text(one.content)
    assert "1. " in one_text and "\n2. " not in one_text
    assert "по мере необходимости Заказчика на основании заявок»." in one_text  # без «.».
    assert "\nЗаказчика" not in one_text  # цитата в одну строку
    empty = client.post(url, headers=auth(api), json={"finding_ids": []})
    assert empty.status_code == 409


def test_clarification_uses_rule_wording(api: dict[str, Any]) -> None:
    client = api["client"]
    body = client.get("/api/rules/VOL-01", headers=auth(api)).json()["body"]
    body["clarification"] = "Просим указать минимальный объём поставки по одной заявке."
    saved = client.put("/api/rules/VOL-01", headers=auth(api, "methodologist"), json={"rule": body})
    assert saved.status_code == 200, saved.text
    created = upload(api)
    response = client.post(
        f"/api/checks/{created['check_id']}/clarification", headers=auth(api), json={}
    )
    assert "Просим указать минимальный объём поставки по одной заявке." in docx_text(
        response.content
    )


def test_disputed_red_finding_marked_and_not_asked_by_default(api: dict[str, Any]) -> None:
    """Второе мнение не согласно — замечание видно, помечено и не уходит в запрос само."""
    api["backend"].disputed = frozenset({"LIA-03"})
    created = upload(api)
    client = api["client"]
    report = client.get(f"/api/checks/{created['check_id']}/report", headers=auth(api)).json()
    penalty = next(f for f in report["findings"] if f["rule_code"] == "LIA-03")
    assert penalty["severity"] == "red" and penalty["needs_review"] is True
    assert penalty["review"] == {"decision": "rejected", "comment": "Цитата о другом."}
    draft = client.post(
        f"/api/checks/{created['check_id']}/clarification", headers=auth(api), json={}
    )
    assert "пеню в размере 1/100" not in docx_text(draft.content)
    word = client.post(f"/api/checks/{created['check_id']}/export?format=docx", headers=auth(api))
    assert "Требует проверки специалистом" in docx_text(word.content)


def test_word_exports_survive_control_characters(api: dict[str, Any]) -> None:
    """Управляющий символ в имени файла или пояснении модели не роняет выгрузку в Word."""
    from sqlalchemy import select

    from svetofor.db.models import Document, Finding

    created = upload(api)
    with api["factory"]() as db:
        for document in db.scalars(select(Document)):
            document.filename = document.filename.replace(".", "\x0b.", 1)
        for finding in db.scalars(select(Finding)):
            finding.explanation = f"{finding.explanation or ''}\x1b[31m\x01"
        db.commit()
    client, check = api["client"], created["check_id"]
    word = client.post(f"/api/checks/{check}/export", headers=auth(api))
    assert word.status_code == 200, word.text
    assert "\x1b" not in docx_text(word.content)
    draft = client.post(f"/api/checks/{check}/clarification", headers=auth(api), json={})
    assert draft.status_code == 200, draft.text
    reports = api["tmp"] / "data"
    assert not [p for p in reports.rglob("tmp*.docx")]  # временные файлы не остаются


def _card(**values: str) -> Any:
    from types import SimpleNamespace

    return SimpleNamespace(card=[SimpleNamespace(key=k, value=v) for k, v in values.items()])


@pytest.mark.parametrize(
    ("card", "expected"),
    [
        ({"law": "44-ФЗ", "application_deadline": "15.10.2026 10:00"}, "12.10.2026"),
        ({"application_deadline": "03.11.2026"}, "31.10.2026"),  # закон не найден — как 44-ФЗ
        ({"law": "223-ФЗ", "application_deadline": "15.10.2026"}, None),  # у 223-ФЗ свои сроки
        ({"law": "44-ФЗ", "application_deadline": "31.02.2026"}, None),  # такой даты нет
        ({"law": "44-ФЗ"}, None),
    ],
)
def test_clarification_request_deadline(card: dict[str, str], expected: str | None) -> None:
    from svetofor.reports.clarification import request_deadline

    assert request_deadline(_card(**card)) == expected


def test_clarification_draft_tells_deadline_and_limit(api: dict[str, Any]) -> None:
    from sqlalchemy import select

    from svetofor.db.models import Chunk, Fact

    created = upload(api)
    with api["factory"]() as db:
        chunk = db.scalars(select(Chunk).limit(1)).one()
        db.add(
            Fact(
                check_id=created["check_id"],
                document_id=chunk.document_id,
                chunk_id=chunk.id,
                quote="Дата и время окончания срока подачи заявок: 15.10.2026 10:00",
                key="application_deadline",
                value_text="15.10.2026 10:00",
                source="rule",
            )
        )
        db.commit()
    client, check = api["client"], created["check_id"]
    report = client.get(f"/api/checks/{check}/report", headers=auth(api)).json()
    assert report["clarification_deadline"] == "12.10.2026"
    draft = client.post(f"/api/checks/{check}/clarification", headers=auth(api), json={})
    text = docx_text(draft.content)
    assert "не позднее 12.10.2026" in text and "не более трёх запросов" in text


@pytest.mark.parametrize(
    ("card", "has_limit"),
    [
        ({"law": "44-ФЗ"}, True),
        ({}, True),  # закон не определён — как 44-ФЗ
        ({"law": "223-ФЗ", "application_deadline": "15.10.2026"}, False),
        ({"law": "commercial"}, False),
    ],
)
def test_request_limit_note_only_for_44fz(card: dict[str, str], has_limit: bool) -> None:
    """Код-ревью 2026-10-08: «не более трёх запросов» (ч. 5 ст. 42 44-ФЗ) показывалось и для
    223-ФЗ, где порядок задаёт положение о закупке."""
    from svetofor.reports.clarification import deadline_note

    note = deadline_note(_card(**card))
    assert (note is not None and "не более трёх запросов" in note) is has_limit
